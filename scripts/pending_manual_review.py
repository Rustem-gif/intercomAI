"""List the chats QA still needs to grade by hand, as an XLSX worklist.

QA's process is to manually re-grade every red/yellow chat in a period. Answering "which ones
did we miss?" by eye is hopeless once a batch has been re-graded, so this asks the database:
graded, below the pass threshold, and no human score yet.

    python scripts/pending_manual_review.py --since 2026-08-01 --until 2026-08-31

Why it exists: the QA rulebook was corrected on 3 and 4 September, which re-graded every chat
with a different score, and several grading runs had crashed mid-batch before that. QA's pass
therefore ran against a list that was still being written, and some red/yellow chats were never
looked at. This reproduces the list as it stands right now.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from intercom_summary.settings import settings
from intercom_summary.storage.db import connect

# Matches the dashboard's colour boundary (frontend lib/utils.ts scoreColor): at or above this
# is green and needs no manual pass; below it is the amber/red band QA re-grades.
DEFAULT_THRESHOLD = 85

_COLS = ["Conversation ID", "Agent", "Date", "AI score", "Band", "Rules version",
         "Graded at", "Summary", "Link"]
_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def rows(conn, since: str, until: str, threshold: int, agent: str | None) -> list[tuple]:
    sql = """
        SELECT c.id, c.agent_name, date(c.created_at) AS day,
               g.overall_score, g.rules_version, substr(g.graded_at,1,16) AS graded_at,
               COALESCE(g.summary,'') AS summary
          FROM conversations c
          JOIN grades g ON g.conversation_id = c.id
         WHERE date(c.created_at) BETWEEN ? AND ?
           AND g.human_score IS NULL
           AND g.overall_score < ?
    """
    args: list[object] = [since, until, threshold]
    if agent:
        sql += " AND c.agent_name = ?"
        args.append(agent)
    sql += " ORDER BY c.agent_name, c.created_at"
    return [tuple(r) for r in conn.execute(sql, args).fetchall()]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--since", required=True, help="YYYY-MM-DD (chat date, inclusive)")
    ap.add_argument("--until", required=True, help="YYYY-MM-DD (chat date, inclusive)")
    ap.add_argument("--threshold", type=int, default=DEFAULT_THRESHOLD)
    ap.add_argument("--agent", help="limit to one agent")
    ap.add_argument("--out", help="XLSX path (default data/exports/pending_manual_review_<range>.xlsx)")
    args = ap.parse_args()

    conn = connect(settings.db_path)
    try:
        data = rows(conn, args.since, args.until, args.threshold, args.agent)
    finally:
        conn.close()

    if not data:
        print("Nothing pending — every chat below the threshold already has a human score.")
        return 0

    by_agent: dict[str, int] = {}
    for r in data:
        by_agent[r[1] or "(unassigned)"] = by_agent.get(r[1] or "(unassigned)", 0) + 1
    print(f"{len(data)} chats below {args.threshold} with no human grade "
          f"({args.since}..{args.until}):")
    for name, n in sorted(by_agent.items(), key=lambda kv: -kv[1]):
        print(f"  {n:4}  {name}")

    out = Path(args.out) if args.out else (
        settings.export_dir / f"pending_manual_review_{args.since}_{args.until}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pending manual review"
    ws.append(_COLS)
    for cid, agent_name, day, score, ver, graded_at, summary in data:
        band = "red" if score < 70 else "yellow"
        ws.append([cid, agent_name, day, score, band, ver, graded_at, summary,
                   f"https://app.intercom.com/a/inbox/_/inbox/conversation/{cid}"])
    for c in range(1, len(_COLS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLS))}1"
    for col, width in zip("ABCDEFGHI", (20, 14, 12, 10, 9, 15, 18, 60, 46)):
        ws.column_dimensions[col].width = width
    wb.save(out)
    print(f"\nWrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
