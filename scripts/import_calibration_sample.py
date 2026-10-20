"""Import a frozen calibration sample from the QA department's .docx list.

    python scripts/import_calibration_sample.py Vybirka_180_Chativ_King_Billy.docx \
        --id kb-v41-180 --name "King Billy v4.1 calibration (180)"

A one-column list — "N. <chat id> — <date> — <agent> — <category>" per line, as in the VIP
sample spreadsheet — is read with --lines (an .xlsx, or plain text):

    python scripts/import_calibration_sample.py "Untitled spreadsheet (1).xlsx" --lines \
        --id vip-v16-70 --name "VIP Scorecard v1.6 calibration (70)" --ruleset vip-v16

The document is the authority on membership, so the import is a straight transcription of its
table — no filtering, no de-duplication against what happens to be in the cache. It then
*reports* which members cannot be graded rather than dropping them, because a short sample is
a finding, not a detail: a member can be absent from the cache, or sitting in the trash, where
it silently blocks its own re-import from Intercom.

Re-running against an existing frozen sample refuses to change it. That is deliberate — the
comparability of two measurement runs is the only thing a calibration set is for.

QA's own verdicts — the answer key the AI is measured against — come in a second pass:

    python scripts/import_calibration_sample.py QA_scores.xlsx --id kb-v41-180 --labels [--dry-run]

`--labels` reads a .docx table, .xlsx or .csv whose header names the columns (chat id, score,
PASS/FAIL, comment — English, Ukrainian or Russian; see LABEL_COLUMNS) and writes them onto the
sample's members. Membership never changes: a chat that is not in the sample is reported.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

from intercom_summary.storage.calibration_store import CalibrationStore, SampleFrozenError

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# The chat id column is the only one we must read exactly; everything else is descriptive.
_ID = re.compile(r"^\d{6,}$")
# The document is Ukrainian; "Так" is the Pilot column's yes.
_YES = {"так", "yes", "y", "+", "true"}


def _cell_text(cell) -> str:
    return " ".join(
        "".join(t.text or "" for t in p.iter(W + "t")) for p in cell.findall(W + "p")
    ).strip()


def rows_from_docx(path: str) -> list[list[str]]:
    root = ET.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    out: list[list[str]] = []
    for table in root.iter(W + "tbl"):
        for row in table.findall(W + "tr"):
            out.append([_cell_text(c) for c in row.findall(W + "tc")])
    return out


# "12. 215561188112687 — 2026-10-07 09:07 — Walter Hofer — Deposit" (dashes may be - – —).
_LINE = re.compile(r"^\s*(\d+)[.)]?\s+(\d{6,})\s*(?:[—–-]\s*(.*))?$")


def rows_from_lines(path: str) -> list[list[str]]:
    """Rows in parse_items' column order from a one-item-per-line list."""
    if path.lower().endswith(".xlsx"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True)
        lines = [str(c) for ws in wb for row in ws.iter_rows(values_only=True)
                 for c in row if c is not None]
    else:
        with open(path, encoding="utf-8") as fh:
            lines = fh.read().splitlines()
    out: list[list[str]] = []
    for line in lines:
        m = _LINE.match(line.strip())
        if not m:
            continue
        rest = [p.strip() for p in re.split(r"\s+[—–-]\s+", m.group(3) or "")] if m.group(3) else []
        out.append([m.group(1), m.group(2), *rest[:3]])
    return out


def parse_items(rows: list[list[str]]) -> list[dict]:
    """Pull the chat rows out of every table in the document.

    Columns are read positionally after the id, which is how the source document is laid out:
    №, Chat ID, date, agent, category, pilot. A row without a numeric chat id in the second
    column is a heading, a progress table or a note, and is skipped.
    """
    items: list[dict] = []
    seen: set[str] = set()
    for r in rows:
        if len(r) < 2 or not _ID.match(r[1]):
            continue
        cid = r[1]
        if cid in seen:            # the document lists each chat once; flag a duplicate loudly
            print(f"  ! duplicate chat id in the document, keeping the first: {cid}")
            continue
        seen.add(cid)
        items.append({
            "conversation_id": cid,
            "seq": int(r[0]) if r[0].isdigit() else len(items) + 1,
            "chat_date": r[2] if len(r) > 2 else None,
            "agent_name": r[3] if len(r) > 3 else None,
            "category": r[4] if len(r) > 4 else None,
            "pilot": len(r) > 5 and r[5].strip().lower() in _YES,
        })
    return items


# Header words → field. Matched case-insensitively as substrings of the header cell, first hit
# wins, so the more specific words come first.
LABEL_COLUMNS: list[tuple[str, tuple[str, ...]]] = [
    ("conversation_id", ("chat id", "chat_id", "conversation id", "conversation_id", "id чат",
                         "id чату", "id чата", "айді", "id")),
    ("human_result", ("pass/fail", "pass / fail", "result", "результат", "вердикт", "verdict")),
    ("human_score", ("qa score", "score", "бал", "оцінка", "оценка", "%")),
    ("human_note", ("comment", "коментар", "комментар", "примітк", "примечан", "note", "reason",
                    "помилк", "ошибк", "feedback")),
]


def rows_from_any(path: str) -> list[list[str]]:
    low = path.lower()
    if low.endswith(".docx"):
        return rows_from_docx(path)
    if low.endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            return [[c.strip() for c in row] for row in csv.reader(fh)]
    if low.endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out: list[list[str]] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append(["" if v is None else str(v).strip() for v in row])
        return out
    raise SystemExit(f"Unsupported file type: {path} (use .docx, .xlsx or .csv)")


def _header_map(row: list[str]) -> dict[str, int]:
    found: dict[str, int] = {}
    for field, words in LABEL_COLUMNS:
        for i, cell in enumerate(row):
            c = cell.lower()
            if i in found.values():
                continue
            if any(w in c for w in words):
                found[field] = i
                break
    return found


def _score(text: str) -> int | None:
    m = re.search(r"\d+(?:[.,]\d+)?", text or "")
    if not m:
        return None
    v = float(m.group().replace(",", "."))
    if 0 < v <= 1 and "." in m.group():          # 0.85 written as a fraction
        v *= 100
    return int(round(v)) if 0 <= v <= 100 else None


def _result(text: str) -> str | None:
    t = (text or "").strip().lower()
    if t.startswith(("pass", "так", "да", "ok", "✅")):
        return "PASS"
    if t.startswith(("fail", "ні", "нет", "❌")):
        return "FAIL"
    return None


def parse_labels(rows: list[list[str]]) -> list[dict]:
    """QA's verdicts from any table with a recognisable header row (re-detected per table, so
    a document with several tables works). A row counts when its id cell is a chat id."""
    labels: list[dict] = []
    cols: dict[str, int] = {}
    for r in rows:
        hm = _header_map(r)
        if "conversation_id" in hm and len(hm) >= 2 and not any(_ID.match(c) for c in r):
            cols = hm
            continue
        if not cols:
            continue
        cid = r[cols["conversation_id"]] if cols["conversation_id"] < len(r) else ""
        cid = cid.split(".")[0]                   # 215561… read from Excel as a float
        if not _ID.match(cid):
            continue
        cell = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else ""  # noqa: E731
        lab = {"conversation_id": cid, "human_score": _score(cell("human_score")),
               "human_result": _result(cell("human_result")), "human_note": cell("human_note")}
        if lab["human_score"] is None and lab["human_result"] is None:
            continue
        labels.append(lab)
    return labels


def import_labels(args) -> int:
    labels = parse_labels(rows_from_any(args.file))
    if not labels:
        print(f"No QA verdicts recognised in {args.file}. The header row must name a chat-id "
              f"column and a score or PASS/FAIL column.", file=sys.stderr)
        return 1
    scored = sum(lab["human_score"] is not None for lab in labels)
    print(f"Parsed {len(labels)} QA verdicts ({scored} with a score) from {args.file}")
    if args.dry_run:
        for lab in labels[:5]:
            print(" ", lab)
        print("  … (dry run, nothing written)")
        return 0
    store = CalibrationStore()
    try:
        if store.get(args.id) is None:
            print(f"No calibration sample {args.id!r} — import its chat list first.", file=sys.stderr)
            return 2
        res = store.set_labels(args.id, labels, labelled_by=args.created_by)
    finally:
        store.close()
    print(f"Labelled {res['labelled']} members of {args.id!r}")
    if res["not_in_sample"]:
        print(f"{len(res['not_in_sample'])} verdicts are for chats NOT in the sample (ignored):")
        for cid in res["not_in_sample"]:
            print(f"  {cid}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("file", help="the .docx list of chats (with --lines, .xlsx / .txt; with --labels, QA's verdicts)")
    ap.add_argument("--lines", action="store_true",
                    help="one 'N. id — date — agent — category' item per line")
    ap.add_argument("--ruleset", help="the sample's default ruleset for runs (e.g. vip-v16)")
    ap.add_argument("--id", required=True, help="sample id, used to scope a review run")
    ap.add_argument("--name", help="sample name (required unless --labels)")
    ap.add_argument("--labels", action="store_true",
                    help="import QA's verdicts onto an existing sample instead of its chat list")
    ap.add_argument("--created-by", default="import script")
    ap.add_argument("--force", action="store_true",
                    help="replace the membership of an already-frozen sample")
    ap.add_argument("--dry-run", action="store_true", help="parse and report, write nothing")
    args = ap.parse_args()
    if args.labels:
        return import_labels(args)
    if not args.name:
        ap.error("--name is required when importing a chat list")

    items = parse_items(rows_from_lines(args.file) if args.lines else rows_from_docx(args.file))
    if not items:
        print(f"No chat rows found in {args.file}", file=sys.stderr)
        return 1
    pilot = sum(1 for i in items if i["pilot"])
    print(f"Parsed {len(items)} chats ({pilot} marked pilot) from {args.file}")

    if args.dry_run:
        for i in items[:5]:
            print(" ", i)
        print("  … (dry run, nothing written)")
        return 0

    store = CalibrationStore()
    try:
        store.create(args.id, args.name, source=args.file, created_by=args.created_by,
                     default_ruleset=args.ruleset)
        try:
            n = store.replace_items(args.id, items, force=args.force)
        except SampleFrozenError as e:
            print(f"{e}", file=sys.stderr)
            return 2
        print(f"Stored {n} chats as calibration sample {args.id!r}")

        missing = store.missing_conversations(args.id)
        trashed = store.trashed_conversations(args.id)
        if missing:
            print(f"\n{len(missing)} of them are in neither the conversations cache nor the "
                  f"VIP threads and cannot be graded until they are fetched:")
            for cid in missing:
                print(f"  {cid}{'   (in the trash)' if cid in set(trashed) else ''}")
        if trashed:
            print(f"\n{len(trashed)} are in the trash. A trashed conversation silently blocks "
                  f"its own re-import from Intercom — restore them before fetching.")
        if not missing and not trashed:
            print("All sample chats are present in the cache.")
    finally:
        store.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
