"""VIP chats tab: the stage-1 dialogue filter, the store, and the API endpoints."""
import textwrap
from datetime import datetime, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from intercom_summary import vip
from intercom_summary.intercom.models import Admin, Contact, Conversation, Message
from intercom_summary.settings import settings
from intercom_summary.storage.vip_store import VipStore


def _convo(messages, cid="1001"):
    return Conversation(
        id=cid, created_at=datetime(2026, 10, 3, tzinfo=timezone.utc), updated_at=None,
        state="closed", subject="VIP", channel="email", brand="Betncare",
        assignee=Admin(id="9", name="Taylor Robinson", email="taylor@co.com"),
        contact=Contact(name="Kylie", email="k@x.com"),
        messages=[Message(i, a, n, None, t, p) for i, (a, n, t, p) in enumerate(messages)],
    )


OFFER = ("admin", "Taylor", "Hi Kylie, are you planning to play today?", "comment")


def test_real_exchange_passes():
    c = _convo([OFFER, ("user", "Kylie", "Yes, I've had a good run. Will play more shortly.", "open"),
                ("admin", "Taylor", "Good luck!", "comment")])
    assert vip.rule_verdict(c) is None


def test_outbound_offer_without_reply_has_no_player():
    assert vip.rule_verdict(_convo([OFFER])) == vip.REASON_NO_PLAYER


@pytest.mark.parametrize("text", ["Thanks!", "Thank you so much 🙏", "Kirrily.A.W.", "ok"])
def test_pleasantries_and_signatures_are_one_word(text):
    c = _convo([OFFER, ("user", "Kylie", text, "open"), ("admin", "Taylor", "Have a great day", "comment")])
    assert vip.rule_verdict(c) == vip.REASON_ONE_WORD


def test_quoted_reply_text_does_not_count_as_the_players_words():
    body = "Thanks\n\nOn Mon, 5 Oct 2026 at 10:00, Taylor <t@co.com> wrote:\n> Are you playing today? Big offer"
    c = _convo([OFFER, ("user", "Kylie", body, "open"), ("admin", "Taylor", "Bye", "comment")])
    assert vip.rule_verdict(c) == vip.REASON_ONE_WORD


def test_request_answered_only_by_a_note_pointing_elsewhere_is_a_duplicate():
    c = _convo([("user", "Jess", "Have I got any free spins please", "comment"),
                ("admin", "Eric", "215561192835356", "note")])
    assert vip.rule_verdict(c) == vip.REASON_DUPLICATE


def test_request_with_only_a_plain_note_is_notes_only():
    c = _convo([("user", "Jess", "Have I got any free spins please", "comment"),
                ("admin", "Eric", "delay, busy", "note")])
    assert vip.rule_verdict(c) == vip.REASON_NOTES_ONLY


def test_system_mail_nobody_answered():
    c = _convo([("user", "Plane", "Yuri has made updates to this work item 99718-50", "comment")])
    assert vip.rule_verdict(c) == vip.REASON_NO_REPLY


def test_window_is_inclusive_utc():
    lo, hi = vip._window("2026-10-01", "2026-10-14")
    assert datetime.fromtimestamp(lo, timezone.utc) == datetime(2026, 10, 1, tzinfo=timezone.utc)
    assert datetime.fromtimestamp(hi, timezone.utc) == datetime(2026, 10, 15, tzinfo=timezone.utc)


def _save(store, cid, keep, stage="ai", reason="ok", topic="Withdrawal"):
    c = _convo([OFFER], cid=cid)
    store.save(vip._row(c, "admin_initiated", vip.counts(c),
                        {"keep": keep, "filter_stage": stage, "reason": reason, "topic": topic}))


def test_store_lists_kept_by_default_and_reports_why_the_rest_went(tmp_path):
    store = VipStore(tmp_path / "vip.db")
    _save(store, "1", 1)
    _save(store, "2", 0, stage="rule", reason=vip.REASON_NO_PLAYER, topic="")
    _save(store, "3", 0, reason="system mail")
    assert [r["id"] for r in store.query()] == ["1"]
    assert len(store.query(include_excluded=True)) == 3
    s = store.summary()
    assert (s["total"], s["kept"], s["excluded"], s["excluded_by_ai"]) == (3, 1, 2, 1)
    assert s["excluded_by_rule"] == [{"filter_stage": "rule", "reason": vip.REASON_NO_PLAYER, "n": 1}]
    assert store.get("1")["payload"]["id"] == "1"
    assert set(store.classified()) == {"1", "2", "3"}
    store.close()


def test_vip_threads_never_reach_the_conversations_table(tmp_path):
    from intercom_summary.storage.conversations_store import ConversationsStore
    store = VipStore(tmp_path / "vip.db")
    _save(store, "1", 1)
    store.close()
    cs = ConversationsStore(tmp_path / "vip.db")
    assert cs.get("1") is None
    cs.close()


@pytest.fixture
def client(tmp_path):
    object.__setattr__(settings, "db_path", tmp_path / "web.db")
    object.__setattr__(settings, "web_basic_auth", "")
    users_file = tmp_path / "web_users.yaml"
    from intercom_summary.web.auth import hash_password
    users_file.write_text(textwrap.dedent(f"""
        users:
          looker:
            password_hash: "{hash_password('pw')}"
            role: viewer
    """))
    from intercom_summary.web import auth as auth_mod
    auth_mod.users = auth_mod.UserStore(users_file)
    store = VipStore(settings.db_path)
    _save(store, "77", 1)
    _save(store, "78", 0, stage="rule", reason=vip.REASON_ONE_WORD, topic="")
    store.close()
    from intercom_summary.web.api import create_app
    c = TestClient(create_app())
    assert c.post("/api/auth/login", json={"username": "looker", "password": "pw"}).status_code == 200
    return c


def test_api_list_detail_and_export(client):
    r = client.get("/api/vip/conversations").json()
    assert [i["id"] for i in r["items"]] == ["77"] and r["summary"]["total"] == 2
    assert len(client.get("/api/vip/conversations?include_excluded=true").json()["items"]) == 2
    d = client.get("/api/vip/conversations/77").json()
    assert d["grade"] is None and "Taylor" in d["transcript"] and d["vip"]["topic"] == "Withdrawal"
    assert client.get("/api/vip/conversations/nope").status_code == 404
    x = client.get("/api/vip/export.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert x.headers["content-disposition"].endswith('filename="vip_threads.xlsx"')

    def sheet(query=""):
        ws = load_workbook(BytesIO(client.get(f"/api/vip/export.xlsx{query}").content)).active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][:2] == ("Conversation ID", "Date (UTC)") and rows[0][-2:] == ("Kept", "Transcript")
        return {r[0]: r for r in rows[1:]}

    kept = sheet()
    assert list(kept) == ["77"] and kept["77"][10] == "yes" and "Taylor" in kept["77"][11]
    # The file follows the view's filters: the dropped threads say why they were dropped.
    both = sheet("?include_excluded=true")
    assert set(both) == {"77", "78"} and both["78"][10] == f"no — rule: {vip.REASON_ONE_WORD}"
    assert sheet("?agent=Nobody") == {} and list(sheet("?topic=Withdrawal")) == ["77"]


def test_collect_query_covers_threads_assigned_to_a_manager_personally():
    """Oct 2026: 617 of the 1,508 emails assigned to a VIP manager carried no team, so a
    team-only search never fetched them."""
    from intercom_summary.vip import VIP_TEAM_IDS, collect_query, vip_member_ids

    admins = [
        {"id": 1, "name": "Alisa Morgan", "team_ids": [VIP_TEAM_IDS[0], 99]},
        {"id": 2, "name": "Support Sam", "team_ids": [99]},
        {"id": 3, "name": "Walter Hofer", "team_ids": [VIP_TEAM_IDS[1]]},
        {"id": 4, "name": "Nobody", "team_ids": None},
    ]
    assert vip_member_ids(admins) == [1, 3]
    q = collect_query(100, 200, [1, 3])
    assigned = q["value"][2]["value"]
    assert {(c["field"], c["value"]) for c in assigned} == {
        ("team_assignee_id", VIP_TEAM_IDS[0]), ("team_assignee_id", VIP_TEAM_IDS[1]),
        ("admin_assignee_id", 1), ("admin_assignee_id", 3),
    }
    assert q["value"][0] == {"field": "created_at", "operator": ">", "value": 99}
