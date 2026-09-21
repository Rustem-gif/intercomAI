from datetime import datetime, timezone
from types import SimpleNamespace

from intercom_summary.intercom.models import Admin, Conversation, Message
from intercom_summary.qa.grader import Grader
from intercom_summary.qa.report import aggregate, report_markdown
from intercom_summary.qa.rules import Ruleset


class FakeAnthropic:
    """Returns a single tool_use block mimicking the Anthropic SDK response."""

    def __init__(self, tool_input):
        self._tool_input = tool_input
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        block = SimpleNamespace(type="tool_use", name="submit_grade", input=self._tool_input)
        return SimpleNamespace(content=[block])


def _convo():
    return Conversation(
        id="42",
        created_at=datetime(2026, 5, 1, tzinfo=timezone.utc),
        updated_at=None,
        state="closed",
        subject="Login",
        assignee=Admin(id="1", name="Ada", email="ada@co.com"),
        messages=[Message(0, "admin", "Ada", None, "Hello, how can I help?")],
    )


def test_grade_parses_tool_output():
    rs = Ruleset(text="rules here", version="abc123", path=None)  # type: ignore[arg-type]
    fake = FakeAnthropic({
        "overall_score": 88,
        "summary": "Handled well.",
        "rule_results": [{"rule_id": "tone-greeting", "verdict": "pass", "evidence": "Hello"}],
        "violations": ["Did not confirm resolution"],
        "suggestions": ["Ask if anything else is needed"],
    })
    grader = Grader(ruleset=rs, model="claude-opus-4-8", client=fake)
    grade = grader.grade(_convo())

    assert grade.overall_score == 88
    assert grade.agent_name == "Ada"
    assert grade.agent_email == "ada@co.com"
    assert grade.rules_version == "abc123"
    assert grade.model == "claude-opus-4-8"
    assert grade.rule_results[0].rule_id == "tone-greeting"
    assert grade.violations == ["Did not confirm resolution"]


def test_aggregate_and_report():
    rs = Ruleset(text="r", version="v", path=None)  # type: ignore[arg-type]
    fake = FakeAnthropic({"overall_score": 70, "summary": "ok",
                          "rule_results": [], "violations": ["X"], "suggestions": []})
    grade = Grader(ruleset=rs, client=fake).grade(_convo())
    agg = aggregate([grade])
    assert agg["Ada"]["count"] == 1
    assert agg["Ada"]["avg_score"] == 70.0
    md = report_markdown([grade])
    assert "Ada" in md and "70/100" in md


def test_report_xlsx_prints_display_names_for_overrides(tmp_path):
    """The export is for humans: "Overridden By" should read the name, not the login."""
    from openpyxl import load_workbook

    from intercom_summary.qa.report import report_xlsx

    rs = Ruleset(text="r", version="v", path=None)  # type: ignore[arg-type]
    fake = FakeAnthropic({"overall_score": 70, "summary": "ok",
                          "rule_results": [], "violations": [], "suggestions": []})
    grade = Grader(ruleset=rs, client=fake).grade(_convo())
    grade.human_score = 90
    grade.overridden_by = "analyst"

    header_and_row = lambda path: load_workbook(path)["Conversations"]

    # Mapped → the name.
    ws = header_and_row(report_xlsx([grade], tmp_path / "named.xlsx", {"analyst": "Daria"}))
    assert ws.cell(row=1, column=5).value == "Overridden By"
    assert ws.cell(row=2, column=5).value == "Daria"

    # Unmapped, and no map at all → the raw username, exactly as before.
    ws = header_and_row(report_xlsx([grade], tmp_path / "other.xlsx", {"kate": "Kate"}))
    assert ws.cell(row=2, column=5).value == "analyst"
    ws = header_and_row(report_xlsx([grade], tmp_path / "bare.xlsx"))
    assert ws.cell(row=2, column=5).value == "analyst"
